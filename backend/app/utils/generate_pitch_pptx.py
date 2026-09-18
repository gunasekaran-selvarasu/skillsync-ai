import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette
    BG_COLOR = RGBColor(248, 250, 252)        # Slate 50
    PRIMARY_PURPLE = RGBColor(124, 58, 237)   # Brand Purple
    DEEP_PURPLE = RGBColor(76, 29, 149)       # Purple 900
    EMERALD_GREEN = RGBColor(16, 185, 129)    # Brand Emerald
    TEXT_DARK = RGBColor(15, 23, 42)          # Slate 900
    TEXT_MUTED = RGBColor(100, 116, 139)      # Slate 500
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)    # Slate 200

    def set_slide_background(slide, color=BG_COLOR):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1, DEEP_PURPLE)

    # Accent decorative shape
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(0.15), Inches(5.5))
    accent.fill.solid()
    accent.fill.fore_color.rgb = EMERALD_GREEN
    accent.line.fill.background()

    # Title box
    txBox = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(3.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "COLLEGE CAREER INTELLIGENCE & PLACEMENT READINESS"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = EMERALD_GREEN
    p_badge.space_after = Pt(14)

    p_title = tf.add_paragraph()
    p_title.text = "SkillSync AI"
    p_title.font.size = Pt(54)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(14)

    p_sub = tf.add_paragraph()
    p_sub.text = "The AI Operating System Connecting Student Capabilities Directly to Campus Placement Eligibility."
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = RGBColor(221, 214, 254)

    p_foot = tf.add_paragraph()
    p_foot.text = "Executive Briefing for College Leadership, Deans, and Placement Directors (TPO)"
    p_foot.font.size = Pt(14)
    p_foot.font.color.rgb = RGBColor(167, 139, 250)
    p_foot.space_before = Pt(36)

    # ==========================================
    # SLIDE 2: The Campus Placement Dilemma
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)

    # Slide Header
    h_box = s2.shapes.add_textbox(Inches(1), Inches(0.7), Inches(11.3), Inches(1))
    htf = h_box.text_frame
    hp = htf.paragraphs[0]
    hp.text = "THE CAMPUS PLACEMENT CRISIS"
    hp.font.size = Pt(12)
    hp.font.bold = True
    hp.font.color.rgb = PRIMARY_PURPLE

    hp2 = htf.add_paragraph()
    hp2.text = "Why 60%+ of College Graduates Struggle in Placement Drives"
    hp2.font.size = Pt(26)
    hp2.font.bold = True
    hp2.font.color.rgb = TEXT_DARK

    # 3 Pain Point Cards
    cards_data = [
        ("For College TPO Leadership", "Placement offices are burdened with manual spreadsheets and guesswork. Basic CGPA filters miss true technical capability, leading to lower recruiter conversion and hiring quotas.", RGBColor(254, 242, 242), RGBColor(239, 68, 68)),
        ("For Faculty Mentors", "Faculty coordinate academic requirements but lack continuous, data-backed visibility into actual student industry competencies until placement season has already begun.", RGBColor(254, 249, 195), RGBColor(202, 138, 4)),
        ("For Graduating Students", "Students submit generic resumes that fail automated corporate ATS filters. They consume random online tutorials without knowing which exact skills unlock high-paying jobs.", RGBColor(243, 232, 255), PRIMARY_PURPLE)
    ]

    for i, (title, desc, fill_col, border_col) in enumerate(cards_data):
        left = 1.0 + (i * 3.9)
        add_card(s2, left, 2.0, 3.6, 4.4, WHITE, border_col)
        
        cbox = s2.shapes.add_textbox(Inches(left + 0.25), Inches(2.3), Inches(3.1), Inches(3.8))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        
        p1 = ctf.paragraphs[0]
        p1.text = f"0{i+1}"
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = border_col
        p1.space_after = Pt(10)
        
        p2 = ctf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        p2.space_after = Pt(12)
        
        p3 = ctf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(13)
        p3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: The Core Principle
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)

    h_box = s3.shapes.add_textbox(Inches(1), Inches(0.7), Inches(11.3), Inches(1))
    htf = h_box.text_frame
    hp = htf.paragraphs[0]
    hp.text = "THE GUIDING PRINCIPLE"
    hp.font.size = Pt(12)
    hp.font.bold = True
    hp.font.color.rgb = PRIMARY_PURPLE

    hp2 = htf.add_paragraph()
    hp2.text = "Transforming Higher Education From Guesswork to Precision"
    hp2.font.size = Pt(26)
    hp2.font.bold = True
    hp2.font.color.rgb = TEXT_DARK

    steps = [
        ("Step 1", "Where Am I Now?", "Continuous AI Career Twin continuously captures verified evidence from assessments, practical challenges, code repositories, and resumes."),
        ("Step 2", "Where Do I Want To Go?", "Real-time benchmark alignments mapping candidates to target careers and active corporate campus openings."),
        ("Step 3", "What Am I Missing?", "Algorithmic skill-gap detection pinpointing exact missing prerequisites and categorizing competencies into Match / Missing."),
        ("Step 4", "What Should I Do Next?", "Singular high-leverage daily action recommendation, dynamic phased learning roadmaps, and gap-closing projects.")
    ]

    for i, (num, headline, body) in enumerate(steps):
        top = 2.0 + (i * 1.25)
        add_card(s3, 1.0, top, 11.33, 1.05, WHITE, PRIMARY_PURPLE if i == 3 else BORDER_COLOR)
        
        tbox = s3.shapes.add_textbox(Inches(1.3), Inches(top + 0.15), Inches(10.7), Inches(0.8))
        tf = tbox.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"{num}: {headline} — "
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_PURPLE if i == 3 else TEXT_DARK
        
        run = p.add_run()
        run.text = body
        run.font.size = Pt(13)
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: Three Portals Architecture
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)

    h_box = s4.shapes.add_textbox(Inches(1), Inches(0.7), Inches(11.3), Inches(1))
    htf = h_box.text_frame
    hp = htf.paragraphs[0]
    hp.text = "INTEGRATED ECOSYSTEM"
    hp.font.size = Pt(12)
    hp.font.bold = True
    hp.font.color.rgb = PRIMARY_PURPLE

    hp2 = htf.add_paragraph()
    hp2.text = "Three Tailored Portals, One Unified Placement Network"
    hp2.font.size = Pt(26)
    hp2.font.bold = True
    hp2.font.color.rgb = TEXT_DARK

    portals = [
        ("STUDENT PORTAL", "Empowering Autonomous Career Growth", [
            "AI Career Twin & Readiness Gauge",
            "Personalized 'What Should I Learn Next?'",
            "Dynamic Roadmap with Task Status Toggling",
            "ATS Resume Studio & Keyword Improver",
            "On-Demand AI Mock Interview Rooms",
            "Career 'What-If' Simulation Sandbox"
        ], PRIMARY_PURPLE),
        ("FACULTY PORTAL", "Proactive Cohort Guidance & Mentorship", [
            "Department Cohort Readiness Monitor",
            "Assessment Completion Tracking",
            "Average Skill Evidence Telemetry",
            "Early At-Risk Candidate Flags",
            "Targeted Mentoring Notes & Action Items",
            "Accreditation-Ready Progress Logs"
        ], RGBColor(37, 99, 235)),
        ("ADMIN / TPO PORTAL", "Automating Enterprise Campus Hiring", [
            "Institutional Placement Intelligence",
            "Department Skill Heatmaps & Trends",
            "1-Click Drive Eligibility Matching",
            "Corporate Partner & Job Management",
            "Auditable Verified Evidence Trail",
            "NAAC / NIRF Data Export"
        ], EMERALD_GREEN)
    ]

    for i, (header, sub, bullets, color) in enumerate(portals):
        left = 1.0 + (i * 3.9)
        add_card(s4, left, 2.0, 3.6, 4.8, WHITE, color)
        
        pbox = s4.shapes.add_textbox(Inches(left + 0.25), Inches(2.2), Inches(3.1), Inches(4.3))
        ptf = pbox.text_frame
        ptf.word_wrap = True
        
        p1 = ptf.paragraphs[0]
        p1.text = header
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = color
        p1.space_after = Pt(4)
        
        p2 = ptf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_after = Pt(14)
        
        for b in bullets:
            bp = ptf.add_paragraph()
            bp.text = f"•  {b}"
            bp.font.size = Pt(12)
            bp.font.color.rgb = TEXT_DARK
            bp.space_after = Pt(6)

    # ==========================================
    # SLIDE 5: Measurable Institutional Impact & ROI
    # ==========================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)

    h_box = s5.shapes.add_textbox(Inches(1), Inches(0.7), Inches(11.3), Inches(1))
    htf = h_box.text_frame
    hp = htf.paragraphs[0]
    hp.text = "PROVEN INSTITUTIONAL VALUE"
    hp.font.size = Pt(12)
    hp.font.bold = True
    hp.font.color.rgb = PRIMARY_PURPLE

    hp2 = htf.add_paragraph()
    hp2.text = "Measurable Outcomes for College Leadership & Admissions"
    hp2.font.size = Pt(26)
    hp2.font.bold = True
    hp2.font.color.rgb = TEXT_DARK

    metrics = [
        ("+35%", "Average Package Growth", "Students present demonstrable, verified projects and ATS-optimized resumes that stand out to tier-1 recruiters."),
        ("80%", "Reduction in TPO Admin Work", "Automatic candidate eligibility matching eliminates hundreds of hours spent manually formatting Excel spreadsheets."),
        ("2x", "Recruiter Retention Rate", "Corporate partners return year after year because candidate pools meet rigorous, pre-verified technical thresholds."),
        ("100%", "Continuous Evidence Audit", "All claimed student competencies are backed by multi-source confidence scores (tests, code, verification).")
    ]

    for i, (stat, title, text) in enumerate(metrics):
        col = i % 2
        row = i // 2
        left = 1.0 + (col * 5.8)
        top = 2.0 + (row * 2.4)
        
        add_card(s5, left, top, 5.5, 2.1, WHITE, BORDER_COLOR)
        
        mbox = s5.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.2), Inches(4.9), Inches(1.7))
        mtf = mbox.text_frame
        mtf.word_wrap = True
        
        p = mtf.paragraphs[0]
        p.text = stat
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_PURPLE if col == 0 else EMERALD_GREEN
        p.space_after = Pt(2)
        
        p2 = mtf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(16)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        p2.space_after = Pt(4)
        
        p3 = mtf.add_paragraph()
        p3.text = text
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 6: Conclusion / Partnership CTA
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6, DEEP_PURPLE)

    cbox = s6.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.33), Inches(4.0))
    ctf = cbox.text_frame
    ctf.word_wrap = True

    cp1 = ctf.paragraphs[0]
    cp1.text = "READY TO TRANSFORM YOUR CAMPUS PLACEMENTS?"
    cp1.font.size = Pt(14)
    cp1.font.bold = True
    cp1.font.color.rgb = EMERALD_GREEN
    cp1.space_after = Pt(16)

    cp2 = ctf.add_paragraph()
    cp2.text = "Empower Your Students, Faculty, and TPO with SkillSync AI"
    cp2.font.size = Pt(38)
    cp2.font.bold = True
    cp2.font.color.rgb = WHITE
    cp2.space_after = Pt(20)

    cp3 = ctf.add_paragraph()
    cp3.text = "Seamless Multi-Tenant Deployment • Zero Disruption to Existing Curriculum • Immediate Student Onboarding"
    cp3.font.size = Pt(16)
    cp3.font.color.rgb = RGBColor(221, 214, 254)
    cp3.space_after = Pt(30)

    cp4 = ctf.add_paragraph()
    cp4.text = "Contact the Institutional Team: partnerships@skillsync.ai | www.skillsync.ai"
    cp4.font.size = Pt(16)
    cp4.font.bold = True
    cp4.font.color.rgb = EMERALD_GREEN

    output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "SkillSync_AI_Pitch.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
